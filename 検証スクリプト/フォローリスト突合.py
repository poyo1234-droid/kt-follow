# -*- coding: utf-8 -*-
"""
キントーンが出力したフォローリストと、旧システムの出力を突き合わせる。

使い方:
    python フォローリスト突合.py            # 出力がある全メーカー
    python フォローリスト突合.py 2          # メーカーNo. で絞る
    python フォローリスト突合.py 2 9 51

突合キーは CS別品番。ヘッダー名が一致する列だけを比較する。
数値は数値として比較する（旧データ側に "1000.0" 形式が混ざるため）。
"""
import sys, re, unicodedata
from pathlib import Path
from collections import defaultdict

import openpyxl

# このスクリプトはリポジトリ直下の 検証スクリプト/ に置く前提
ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "0907フォローデータ提出用" / "0907フォローデータ提出用"
KINTONE_DIR = BASE / "フォローリスト-kintone出力"
KYUU_DIR    = BASE / "フォローリスト20260907"

KEY = "CS別品番"


def maker_no(name: str):
    """ファイル名の先頭の数字をメーカーNo. として取り出す（00002 も 2 も '2' にする）"""
    m = re.match(r"^(\d+)", name)
    return str(int(m.group(1))) if m else None


def load(path: Path):
    """xlsx を読み、{CS別品番: {列名: 値}} と列名リストを返す"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    idx = {h: i for i, h in enumerate(hdr)}
    if KEY not in idx:
        wb.close()
        raise SystemExit(f"[エラー] {path.name} に「{KEY}」列がありません")
    rows = {}
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in row):
            continue
        k = str(row[idx[KEY]]).strip() if row[idx[KEY]] is not None else ""
        if not k:
            continue
        if k not in rows:                       # 重複は最初のものを採る
            rows[k] = {h: row[i] for h, i in idx.items()}
    wb.close()
    return hdr, rows


def collect(dir_path: Path):
    """メーカーNo. → [ファイル] の辞書。旧側は子カテゴリ単位で複数ファイルに分かれることがある"""
    out = defaultdict(list)
    for p in sorted(dir_path.glob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        no = maker_no(p.name)
        if no:
            out[no].append(p)
    return out


def norm(v):
    """比較用に正規化する。数値として読めるものは float にそろえる"""
    if v is None:
        return ""
    if isinstance(v, str):
        v = unicodedata.normalize("NFKC", v).strip()
        if v == "":
            return ""
    try:
        return float(v)
    except (TypeError, ValueError):
        return str(v).strip()


def compare(no, k_files, o_files):
    k_hdr, k_rows = {}, {}
    for p in k_files:
        h, r = load(p)
        k_hdr = h
        k_rows.update(r)
    o_hdr, o_rows = {}, {}
    for p in o_files:
        h, r = load(p)
        o_hdr = h
        o_rows.update(r)

    cols = [h for h in k_hdr if h in o_hdr and h is not None]
    only_k = set(k_rows) - set(o_rows)
    only_o = set(o_rows) - set(k_rows)
    common = set(k_rows) & set(o_rows)

    print(f"\n{'='*70}")
    print(f"メーカーNo. {no}")
    print(f"  キントーン {len(k_rows):>6,} 件 ({len(k_files)}ファイル) / "
          f"旧システム {len(o_rows):>6,} 件 ({len(o_files)}ファイル)")
    print(f"  共通 {len(common):,} / キントーンのみ {len(only_k):,} / 旧のみ {len(only_o):,}")
    print(f"  比較列 {len(cols)} 列")

    mismatch = {c: 0 for c in cols}
    example = {c: [] for c in cols}
    for cs in common:
        kr, orow = k_rows[cs], o_rows[cs]
        for c in cols:
            kv, ov = norm(kr.get(c)), norm(orow.get(c))
            if kv != ov:
                mismatch[c] += 1
                if len(example[c]) < 2:
                    example[c].append((cs, orow.get(c), kr.get(c)))

    print(f"\n  {'列':<26}{'不一致':>8}")
    print("  " + "-" * 60)
    for c in cols:
        label = str(c).replace("\n", "")
        n = mismatch[c]
        mark = "" if n == 0 else "  ←"
        print(f"  {label:<26}{n:>8,}{mark}")
        for cs, ov, kv in example[c]:
            print(f"      例 {cs}  旧={ov!r}  kintone={kv!r}")
    return len(only_k), len(only_o), sum(mismatch.values())


def main():
    if not KINTONE_DIR.is_dir() or not KYUU_DIR.is_dir():
        raise SystemExit(f"[エラー] 対象フォルダが見つかりません\n  {KINTONE_DIR}\n  {KYUU_DIR}")

    kin, kyu = collect(KINTONE_DIR), collect(KYUU_DIR)
    targets = sys.argv[1:] or sorted(kin.keys(), key=int)

    total = 0
    for no in targets:
        no = str(int(no))
        if no not in kin:
            print(f"[スキップ] メーカー{no}: キントーン側の出力がありません")
            continue
        if no not in kyu:
            print(f"[スキップ] メーカー{no}: 旧システム側の出力がありません")
            continue
        _, _, mm = compare(no, kin[no], kyu[no])
        total += mm

    print(f"\n{'='*70}")
    print(f"不一致セルの合計: {total:,}")


if __name__ == "__main__":
    main()
